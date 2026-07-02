#!/usr/bin/env python3
"""Builds the ICAI Guidance Note (Division I) financial statement workbook.

Formula-only (no macros) so it works in any Excel/LibreOffice with zero
macro security prompts. VBA add-ons (Tally import, nav buttons) are shipped
separately as importable .bas source since this environment has no Excel to
compile/verify a real VBA project.
"""
import os
import sys
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, SCRIPT_DIR)
from groups import GROUPS

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side, NamedStyle
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import FormulaRule

TB_MAX_ROW = 504  # data rows 5..504 => 500 ledger rows capacity
TB_FIRST_ROW = 5

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SECTION_FONT = Font(bold=True, italic=True, color="374151")
TOTAL_FONT = Font(bold=True)
TITLE_FONT = Font(bold=True, size=14)
SUB_FONT = Font(size=10, color="6B7280")
THIN = Side(style="thin", color="D1D5DB")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

wb = Workbook()
wb.remove(wb.active)

# ---------------------------------------------------------------- Master --
master = wb.create_sheet("Master")

master["A1"] = "Key"
master["B1"] = "Label"
master["C1"] = "Statement"
master["D1"] = "Side"
master["E1"] = "Section"
master["F1"] = "Order"
master["G1"] = "EntityTypes"
master["H1"] = "UsedFlag"
master["I1"] = "NoteNo"
for c in range(1, 10):
    master.cell(row=1, column=c).font = Font(bold=True)

group_row_of = {}
for i, (key, label, statement, side, section, order, etypes, keywords) in enumerate(GROUPS):
    r = i + 2
    group_row_of[key] = r
    master.cell(row=r, column=1, value=key)
    master.cell(row=r, column=2, value=label)
    master.cell(row=r, column=3, value=statement)
    master.cell(row=r, column=4, value=side)
    master.cell(row=r, column=5, value=section)
    master.cell(row=r, column=6, value=order)
    master.cell(row=r, column=7, value=",".join(etypes) if etypes else "")
    # UsedFlag: does TB have any row classified into this label?
    master.cell(row=r, column=8,
                value=f'=IF(COUNTIF(TB!$E${TB_FIRST_ROW}:$E${TB_MAX_ROW},B{r})>0,1,0)')

n_groups = len(GROUPS)
master_last_row = n_groups + 1

# NoteNo: sequential count of groups with Order <= this row's Order AND UsedFlag=1
# (mirrors engine.ts assignNoteNumbers - only groups actually used get a number)
for i in range(n_groups):
    r = i + 2
    master.cell(row=r, column=9,
                value=(f'=IF(H{r}=1,SUMPRODUCT(($F$2:$F${master_last_row}<=F{r})'
                       f'*($H$2:$H${master_last_row}=1)),"")'))

# Keyword table (one keyword per row, in priority order == GROUPS order so the
# first / most specific match wins)
master["K1"] = "Keyword"
master["L1"] = "GroupLabel"
master["K1"].font = Font(bold=True)
master["L1"].font = Font(bold=True)
kw_row = 2
for key, label, statement, side, section, order, etypes, keywords in GROUPS:
    for kw in keywords:
        master.cell(row=kw_row, column=11, value=kw)
        master.cell(row=kw_row, column=12, value=label)
        kw_row += 1
kw_last_row = kw_row - 1

master.column_dimensions["A"].width = 22
master.column_dimensions["B"].width = 55
master.column_dimensions["K"].width = 26
master.column_dimensions["L"].width = 45

# Named ranges
wb.defined_names["GroupLabels"] = DefinedName("GroupLabels", attr_text=f"Master!$B$2:$B${master_last_row}")
wb.defined_names["KeywordRange"] = DefinedName("KeywordRange", attr_text=f"Master!$K$2:$K${kw_last_row}")
wb.defined_names["KeywordGroupRange"] = DefinedName("KeywordGroupRange", attr_text=f"Master!$L$2:$L${kw_last_row}")

print("Master sheet built:", n_groups, "groups,", kw_last_row - 1, "keywords")

# ------------------------------------------------------------------ Home --
home = wb.create_sheet("Home")
home.sheet_view.showGridLines = False
home.column_dimensions["A"].width = 3
home.column_dimensions["B"].width = 28
home.column_dimensions["C"].width = 32
home.column_dimensions["D"].width = 3
home.column_dimensions["E"].width = 28
home.column_dimensions["F"].width = 32

home["B2"] = "Financial Statement Preparation Tool"
home["B2"].font = TITLE_FONT
home["B3"] = "ICAI Guidance Note on Financial Statements of Non-Corporate Entities (Division I)"
home["B3"].font = SUB_FONT

home["B5"] = "1. Fill in entity details below."
home["B6"] = "2. Go to the Trial Balance sheet, paste/enter your ledgers and confirm the Group for each."
home["B7"] = "3. Balance Sheet, Profit & Loss and Notes sheets update automatically."
for r in range(5, 8):
    home.cell(row=r, column=2).font = Font(size=10, color="374151")

fields_left = [
    ("Entity Name", "EntityName", "[Enter Entity Name]"),
    ("Entity Type", "EntityType", "Partnership Firm"),
    ("PAN", "EntityPAN", ""),
    ("Nature of Business", "NatureOfBusiness", ""),
    ("Registered Address", "EntityAddress", ""),
]
fields_right = [
    ("Financial Year Ended", "FYEnd", None),
    ("Previous Year Ended", "PYEnd", None),
    ("Rounding", "Rounding", "Actual"),
    ("Has own Income-tax liability? (Y/N)", "HasTax", "Y"),
    ("Show Partners' Remuneration line? (Y/N)", "ShowRemun", "Y"),
]

row0 = 10
home.cell(row=row0 - 1, column=2, value="Entity Details").font = Font(bold=True)
for i, (label, name, default) in enumerate(fields_left):
    r = row0 + i
    home.cell(row=r, column=2, value=label)
    cell = home.cell(row=r, column=3, value=default)
    cell.border = BOX
    wb.defined_names[name] = DefinedName(name, attr_text=f"Home!$C${r}")

home.cell(row=row0 - 1, column=5, value="Reporting").font = Font(bold=True)
for i, (label, name, default) in enumerate(fields_right):
    r = row0 + i
    home.cell(row=r, column=5, value=label)
    cell = home.cell(row=r, column=6, value=default)
    cell.border = BOX
    wb.defined_names[name] = DefinedName(name, attr_text=f"Home!$F${r}")

dv_entity_type = DataValidation(type="list", formula1='"Sole Proprietorship,Partnership Firm,LLP"', allow_blank=False)
home.add_data_validation(dv_entity_type)
dv_entity_type.add(home["C11"])

dv_rounding = DataValidation(type="list", formula1='"Actual,Hundreds,Thousands,Lakhs,Crores,Millions"', allow_blank=False)
home.add_data_validation(dv_rounding)
dv_rounding.add(home["F12"])

dv_yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=False)
home.add_data_validation(dv_yn)
dv_yn.add(home["F13"])
dv_yn2 = DataValidation(type="list", formula1='"Y,N"', allow_blank=False)
home.add_data_validation(dv_yn2)
dv_yn2.add(home["F14"])

signrow = row0 + len(fields_left) + 3
home.cell(row=signrow - 1, column=2, value="Signatory / Preparer").font = Font(bold=True)
sign_fields = [
    ("Name of Signatory", "SignatoryName", ""),
    ("Designation", "SignatoryDesignation", "Partner"),
    ("Place", "SignatoryPlace", ""),
    ("Date", "SignatoryDate", None),
    ("Auditor / Preparer Firm Name", "PreparerFirm", ""),
    ("Membership No.", "PreparerMembership", ""),
]
for i, (label, name, default) in enumerate(sign_fields):
    r = signrow + i
    home.cell(row=r, column=2, value=label)
    cell = home.cell(row=r, column=3, value=default)
    cell.border = BOX
    wb.defined_names[name] = DefinedName(name, attr_text=f"Home!$C${r}")

navrow = signrow + len(sign_fields) + 3
home.cell(row=navrow - 1, column=2, value="Go to").font = Font(bold=True)
for i, sheet_name in enumerate(["Trial Balance -> 'TB'", "Balance Sheet -> 'BS'", "Profit & Loss -> 'PL'", "Notes to Accounts -> 'Notes'"]):
    home.cell(row=navrow + i, column=2, value=sheet_name).font = Font(color="2563EB", underline="single")

wb.save(os.path.join(SCRIPT_DIR, "_wip.xlsx"))
print("Home sheet built")

# -------------------------------------------------------------------- TB --
tb = wb.create_sheet("TB")
tb.freeze_panes = "A5"
widths = {"A": 10, "B": 32, "C": 20, "D": 34, "E": 34, "F": 14, "G": 14, "H": 14, "I": 14, "J": 22}
for col, w in widths.items():
    tb.column_dimensions[col].width = w

tb["A1"] = "=Home!$C$10"
tb["A1"].font = TITLE_FONT
tb["A2"] = "Trial Balance"
tb["A2"].font = Font(bold=True)

headers = ["GL Code", "Ledger Name", "Tally Group (optional)", "Suggested Group", "Group (confirm/override)",
           "CY Debit", "CY Credit", "PY Debit", "PY Credit", "Narration"]
for i, h in enumerate(headers):
    c = tb.cell(row=4, column=i + 1, value=h)
    c.font = HEADER_FONT
    c.fill = HEADER_FILL
    c.alignment = Alignment(wrap_text=True, vertical="center")

dv_group = DataValidation(type="list", formula1="=GroupLabels", allow_blank=True)
tb.add_data_validation(dv_group)

unclassified_fill = PatternFill("solid", fgColor="FEF3C7")
tb.conditional_formatting.add(
    f"B{TB_FIRST_ROW}:E{TB_MAX_ROW}",
    FormulaRule(formula=[f"AND($B{TB_FIRST_ROW}<>\"\",$E{TB_FIRST_ROW}=\"\")"], fill=unclassified_fill),
)

for r in range(TB_FIRST_ROW, TB_MAX_ROW + 1):
    d_cell = f"D{r}"
    tb.cell(row=r, column=4,
            value=ArrayFormula(
                d_cell,
                f'=IF(B{r}="","",IFERROR(INDEX(KeywordGroupRange,'
                f'MATCH(TRUE,ISNUMBER(SEARCH(KeywordRange,B{r}&" "&C{r})),0)),""))'
            ))
    tb.cell(row=r, column=5, value=f'=IF(B{r}="","",D{r})')
    dv_group.add(tb.cell(row=r, column=5))
    for col in (6, 7, 8, 9):
        tb.cell(row=r, column=col).number_format = "#,##0.00;(#,##0.00);-"

total_row = TB_MAX_ROW + 1
tb.cell(row=total_row, column=2, value="TOTAL").font = TOTAL_FONT
for col in (6, 7, 8, 9):
    letter = get_column_letter(col)
    cell = tb.cell(row=total_row, column=col, value=f"=SUM({letter}{TB_FIRST_ROW}:{letter}{TB_MAX_ROW})")
    cell.font = TOTAL_FONT
    cell.number_format = "#,##0.00;(#,##0.00);-"

tb.cell(row=total_row + 2, column=2, value="CY Debit - CY Credit (informational; check the Balance Sheet tie-out check for classification errors)")
tb.cell(row=total_row + 2, column=6, value=f"=F{total_row}-G{total_row}").number_format = "#,##0.00;(#,##0.00);-"
tb.cell(row=total_row + 3, column=2, value="PY Debit - PY Credit (informational)")
tb.cell(row=total_row + 3, column=6, value=f"=H{total_row}-I{total_row}").number_format = "#,##0.00;(#,##0.00);-"

wb.save(os.path.join(SCRIPT_DIR, "_wip.xlsx"))
print("TB sheet built")

# Rounding factor helper, referenced by BS/PL/Notes amount formulas.
master["N1"] = "RoundingFactor"
master["O1"] = ('=CHOOSE(MATCH(Home!$F$12,{"Actual","Hundreds","Thousands","Lakhs","Crores","Millions"},0),'
                 '1,100,1000,100000,10000000,1000000)')
wb.defined_names["RoundingFactor"] = DefinedName("RoundingFactor", attr_text="Master!$O$1")

TB_GROUP_RANGE = f"TB!$E${TB_FIRST_ROW}:$E${TB_MAX_ROW}"
TB_CYD = f"TB!$F${TB_FIRST_ROW}:$F${TB_MAX_ROW}"
TB_CYC = f"TB!$G${TB_FIRST_ROW}:$G${TB_MAX_ROW}"
TB_PYD = f"TB!$H${TB_FIRST_ROW}:$H${TB_MAX_ROW}"
TB_PYC = f"TB!$I${TB_FIRST_ROW}:$I${TB_MAX_ROW}"


def amount_formula(label_ref, side, period):
    """period: 'CY' or 'PY'. side: EquityLiability/Income (credit-natured) or Asset/Expense (debit-natured)."""
    d, c = (TB_CYD, TB_CYC) if period == "CY" else (TB_PYD, TB_PYC)
    if side in ("EquityLiability", "Income"):
        raw = f"(SUMIFS({c},{TB_GROUP_RANGE},{label_ref})-SUMIFS({d},{TB_GROUP_RANGE},{label_ref}))"
    else:
        raw = f"(SUMIFS({d},{TB_GROUP_RANGE},{label_ref})-SUMIFS({c},{TB_GROUP_RANGE},{label_ref}))"
    return f"={raw}/RoundingFactor"


def note_formula(label_ref):
    return f'=IFERROR(INDEX(Master!$I$2:$I${master_last_row},MATCH({label_ref},Master!$B$2:$B${master_last_row},0)),"")'


def group_lookup(section, statement):
    return sorted(
        [g for g in GROUPS if g[2] == statement and g[4] == section],
        key=lambda g: g[5],
    )


BS_EL_SECTIONS = ["Owners' Fund", "Non-current liabilities", "Current liabilities"]
BS_ASSET_SECTIONS = ["Non-current assets", "Current assets"]

bs = wb.create_sheet("BS")
bs.freeze_panes = "A5"
for col, w in {"A": 3, "B": 50, "C": 8, "D": 18, "E": 18}.items():
    bs.column_dimensions[col].width = w

bs["B1"] = "=Home!$C$10"
bs["B1"].font = TITLE_FONT
bs["B2"] = '="Balance Sheet as at "&TEXT(Home!$F$10,"DD MMMM YYYY")'
bs["B2"].font = Font(bold=True)
bs["B3"] = '="(Rs. in "&Home!$F$12&")"'
bs["B3"].font = SUB_FONT

hdr_row = 5
for i, h in enumerate(["Particulars", "Note", f'=TEXT(Home!$F$10,"DD MMM YYYY")', f'=TEXT(Home!$F$11,"DD MMM YYYY")']):
    c = bs.cell(row=hdr_row, column=i + 2, value=h)
    c.font = HEADER_FONT
    c.fill = HEADER_FILL

r = hdr_row + 1
bs.cell(row=r, column=2, value="I. EQUITY AND LIABILITIES").font = TOTAL_FONT
r += 1
el_subtotal_cells = []
for section in BS_EL_SECTIONS:
    bs.cell(row=r, column=2, value=section).font = SECTION_FONT
    r += 1
    line_start = r
    for key, label, statement, side, sect, order, etypes, keywords in group_lookup(section, "BS"):
        label_ref = f'Master!$B${group_row_of[key]}'
        bs.cell(row=r, column=2, value=f"={label_ref}")
        bs.cell(row=r, column=3, value=note_formula(label_ref))
        bs.cell(row=r, column=4, value=amount_formula(label_ref, side, "CY")).number_format = "#,##0.00;(#,##0.00);-"
        bs.cell(row=r, column=5, value=amount_formula(label_ref, side, "PY")).number_format = "#,##0.00;(#,##0.00);-"
        r += 1
    line_end = r - 1
    bs.cell(row=r, column=2, value=f"Total {section}").font = TOTAL_FONT
    if line_end >= line_start:
        bs.cell(row=r, column=4, value=f"=SUM(D{line_start}:D{line_end})").font = TOTAL_FONT
        bs.cell(row=r, column=5, value=f"=SUM(E{line_start}:E{line_end})").font = TOTAL_FONT
    bs.cell(row=r, column=4).number_format = "#,##0.00;(#,##0.00);-"
    bs.cell(row=r, column=5).number_format = "#,##0.00;(#,##0.00);-"
    el_subtotal_cells.append(r)
    r += 2

total_el_row = r
bs.cell(row=r, column=2, value="TOTAL EQUITY AND LIABILITIES").font = TOTAL_FONT
bs.cell(row=r, column=4, value="=" + "+".join(f"D{x}" for x in el_subtotal_cells)).font = TOTAL_FONT
bs.cell(row=r, column=5, value="=" + "+".join(f"E{x}" for x in el_subtotal_cells)).font = TOTAL_FONT
bs.cell(row=r, column=4).number_format = "#,##0.00;(#,##0.00);-"
bs.cell(row=r, column=5).number_format = "#,##0.00;(#,##0.00);-"
r += 2

bs.cell(row=r, column=2, value="II. ASSETS").font = TOTAL_FONT
r += 1
asset_subtotal_cells = []
for section in BS_ASSET_SECTIONS:
    bs.cell(row=r, column=2, value=section).font = SECTION_FONT
    r += 1
    line_start = r
    for key, label, statement, side, sect, order, etypes, keywords in group_lookup(section, "BS"):
        label_ref = f'Master!$B${group_row_of[key]}'
        bs.cell(row=r, column=2, value=f"={label_ref}")
        bs.cell(row=r, column=3, value=note_formula(label_ref))
        bs.cell(row=r, column=4, value=amount_formula(label_ref, side, "CY")).number_format = "#,##0.00;(#,##0.00);-"
        bs.cell(row=r, column=5, value=amount_formula(label_ref, side, "PY")).number_format = "#,##0.00;(#,##0.00);-"
        r += 1
    line_end = r - 1
    bs.cell(row=r, column=2, value=f"Total {section}").font = TOTAL_FONT
    if line_end >= line_start:
        bs.cell(row=r, column=4, value=f"=SUM(D{line_start}:D{line_end})").font = TOTAL_FONT
        bs.cell(row=r, column=5, value=f"=SUM(E{line_start}:E{line_end})").font = TOTAL_FONT
    bs.cell(row=r, column=4).number_format = "#,##0.00;(#,##0.00);-"
    bs.cell(row=r, column=5).number_format = "#,##0.00;(#,##0.00);-"
    asset_subtotal_cells.append(r)
    r += 2

total_assets_row = r
bs.cell(row=r, column=2, value="TOTAL ASSETS").font = TOTAL_FONT
bs.cell(row=r, column=4, value="=" + "+".join(f"D{x}" for x in asset_subtotal_cells)).font = TOTAL_FONT
bs.cell(row=r, column=5, value="=" + "+".join(f"E{x}" for x in asset_subtotal_cells)).font = TOTAL_FONT
bs.cell(row=r, column=4).number_format = "#,##0.00;(#,##0.00);-"
bs.cell(row=r, column=5).number_format = "#,##0.00;(#,##0.00);-"
r += 2

bs.cell(row=r, column=2, value="Tie-out check (Total Equity & Liabilities - Total Assets)").font = Font(italic=True)
bs.cell(row=r, column=4, value=f"=D{total_el_row}-D{total_assets_row}").number_format = "#,##0.00;(#,##0.00);-"
bs.cell(row=r, column=5, value=f"=E{total_el_row}-E{total_assets_row}").number_format = "#,##0.00;(#,##0.00);-"

BS_TOTAL_EL_ROW = total_el_row
BS_TOTAL_ASSETS_ROW = total_assets_row

wb.save(os.path.join(SCRIPT_DIR, "_wip.xlsx"))
print("BS sheet built, tie-out row", r)

# -------------------------------------------------------------------- PL --
pl = wb.create_sheet("PL")
pl.freeze_panes = "A5"
for col, w in {"A": 3, "B": 50, "C": 8, "D": 18, "E": 18}.items():
    pl.column_dimensions[col].width = w

pl["B1"] = "=Home!$C$10"
pl["B1"].font = TITLE_FONT
pl["B2"] = '="Statement of Profit and Loss for the year ended "&TEXT(Home!$F$10,"DD MMMM YYYY")'
pl["B2"].font = Font(bold=True)
pl["B3"] = '="(Rs. in "&Home!$F$12&")"'
pl["B3"].font = SUB_FONT

hdr_row = 5
for i, h in enumerate(["Particulars", "Note", '=TEXT(Home!$F$10,"DD MMM YYYY")', '=TEXT(Home!$F$11,"DD MMM YYYY")']):
    c = pl.cell(row=hdr_row, column=i + 2, value=h)
    c.font = HEADER_FONT
    c.fill = HEADER_FILL

r = hdr_row + 1
pl.cell(row=r, column=2, value="INCOME").font = TOTAL_FONT
r += 1
income_start = r
for key, label, statement, side, sect, order, etypes, keywords in group_lookup("Income", "PL"):
    label_ref = f'Master!$B${group_row_of[key]}'
    pl.cell(row=r, column=2, value=f"={label_ref}")
    pl.cell(row=r, column=3, value=note_formula(label_ref))
    pl.cell(row=r, column=4, value=amount_formula(label_ref, side, "CY")).number_format = "#,##0.00;(#,##0.00);-"
    pl.cell(row=r, column=5, value=amount_formula(label_ref, side, "PY")).number_format = "#,##0.00;(#,##0.00);-"
    r += 1
income_end = r - 1
total_income_row = r
pl.cell(row=r, column=2, value="Total Income").font = TOTAL_FONT
pl.cell(row=r, column=4, value=f"=SUM(D{income_start}:D{income_end})").font = TOTAL_FONT
pl.cell(row=r, column=5, value=f"=SUM(E{income_start}:E{income_end})").font = TOTAL_FONT
pl.cell(row=r, column=4).number_format = "#,##0.00;(#,##0.00);-"
pl.cell(row=r, column=5).number_format = "#,##0.00;(#,##0.00);-"
r += 2

pl.cell(row=r, column=2, value="EXPENSES").font = TOTAL_FONT
r += 1
expense_start = r
for key, label, statement, side, sect, order, etypes, keywords in group_lookup("Expenses", "PL"):
    label_ref = f'Master!$B${group_row_of[key]}'
    pl.cell(row=r, column=2, value=f"={label_ref}")
    pl.cell(row=r, column=3, value=note_formula(label_ref))
    pl.cell(row=r, column=4, value=amount_formula(label_ref, side, "CY")).number_format = "#,##0.00;(#,##0.00);-"
    pl.cell(row=r, column=5, value=amount_formula(label_ref, side, "PY")).number_format = "#,##0.00;(#,##0.00);-"
    r += 1
expense_end = r - 1
total_expense_row = r
pl.cell(row=r, column=2, value="Total Expenses").font = TOTAL_FONT
pl.cell(row=r, column=4, value=f"=SUM(D{expense_start}:D{expense_end})").font = TOTAL_FONT
pl.cell(row=r, column=5, value=f"=SUM(E{expense_start}:E{expense_end})").font = TOTAL_FONT
pl.cell(row=r, column=4).number_format = "#,##0.00;(#,##0.00);-"
pl.cell(row=r, column=5).number_format = "#,##0.00;(#,##0.00);-"
r += 2

before_remun_row = r
pl.cell(row=r, column=2,
        value='=IF(AND(Home!$C$11="Partnership Firm",Home!$F$14="Y"),'
              '"Profit before partners\' remuneration and tax","Profit before tax")').font = TOTAL_FONT
pl.cell(row=r, column=4, value=f"=D{total_income_row}-D{total_expense_row}").font = TOTAL_FONT
pl.cell(row=r, column=5, value=f"=E{total_income_row}-E{total_expense_row}").font = TOTAL_FONT
pl.cell(row=r, column=4).number_format = "#,##0.00;(#,##0.00);-"
pl.cell(row=r, column=5).number_format = "#,##0.00;(#,##0.00);-"
r += 1

remun_key = "partners_remuneration"
remun_label_ref = f'Master!$B${group_row_of[remun_key]}'
remun_row = r
pl.cell(row=r, column=2,
        value='=IF(AND(Home!$C$11="Partnership Firm",Home!$F$14="Y"),"Less: Remuneration and interest to partners","")')
pl.cell(r, 4,
        value=f'=IF(AND(Home!$C$11="Partnership Firm",Home!$F$14="Y"),'
              + amount_formula(remun_label_ref, "Expense", "CY")[1:] + ',0)').number_format = "#,##0.00;(#,##0.00);-"
pl.cell(r, 5,
        value=f'=IF(AND(Home!$C$11="Partnership Firm",Home!$F$14="Y"),'
              + amount_formula(remun_label_ref, "Expense", "PY")[1:] + ',0)').number_format = "#,##0.00;(#,##0.00);-"
r += 1

pbt_row = r
pl.cell(row=r, column=2, value="Profit before tax").font = TOTAL_FONT
pl.cell(row=r, column=4, value=f"=D{before_remun_row}-D{remun_row}").font = TOTAL_FONT
pl.cell(row=r, column=5, value=f"=E{before_remun_row}-E{remun_row}").font = TOTAL_FONT
pl.cell(row=r, column=4).number_format = "#,##0.00;(#,##0.00);-"
pl.cell(row=r, column=5).number_format = "#,##0.00;(#,##0.00);-"
r += 2

pl.cell(row=r, column=2, value='=IF(Home!$F$13="Y","TAX EXPENSE","")').font = TOTAL_FONT
r += 1
ctax_ref = f'Master!$B${group_row_of["current_tax"]}'
dtax_ref = f'Master!$B${group_row_of["deferred_tax_pl"]}'
ctax_row = r
pl.cell(row=r, column=2, value='=IF(Home!$F$13="Y","Current tax","")')
pl.cell(r, 4, value=f'=IF(Home!$F$13="Y",' + amount_formula(ctax_ref, "Expense", "CY")[1:] + ',0)').number_format = "#,##0.00;(#,##0.00);-"
pl.cell(r, 5, value=f'=IF(Home!$F$13="Y",' + amount_formula(ctax_ref, "Expense", "PY")[1:] + ',0)').number_format = "#,##0.00;(#,##0.00);-"
r += 1
dtax_row = r
pl.cell(row=r, column=2, value='=IF(Home!$F$13="Y","Deferred tax","")')
pl.cell(r, 4, value=f'=IF(Home!$F$13="Y",' + amount_formula(dtax_ref, "Expense", "CY")[1:] + ',0)').number_format = "#,##0.00;(#,##0.00);-"
pl.cell(r, 5, value=f'=IF(Home!$F$13="Y",' + amount_formula(dtax_ref, "Expense", "PY")[1:] + ',0)').number_format = "#,##0.00;(#,##0.00);-"
r += 2

pl.cell(row=r, column=2, value="PROFIT FOR THE YEAR").font = TOTAL_FONT
pl.cell(row=r, column=4, value=f"=D{pbt_row}-D{ctax_row}-D{dtax_row}").font = TOTAL_FONT
pl.cell(row=r, column=5, value=f"=E{pbt_row}-E{ctax_row}-E{dtax_row}").font = TOTAL_FONT
pl.cell(row=r, column=4).number_format = "#,##0.00;(#,##0.00);-"
pl.cell(row=r, column=5).number_format = "#,##0.00;(#,##0.00);-"

wb.save(os.path.join(SCRIPT_DIR, "_wip.xlsx"))
print("PL sheet built, profit row", r)

# ----------------------------------------------------------------- Notes --
notes = wb.create_sheet("Notes")
for col, w in {"A": 3, "B": 50, "C": 18, "D": 18}.items():
    notes.column_dimensions[col].width = w

notes["B1"] = "=Home!$C$10"
notes["B1"].font = TITLE_FONT
notes["B2"] = "Notes forming part of the Financial Statements"
notes["B2"].font = Font(bold=True)
notes["B3"] = ('="Tip: filter the Group column on the Trial Balance sheet by a Note\'s heading below to see the '
               'ledger-level detail behind that note."')
notes["B3"].font = SUB_FONT

r = 5
for g in sorted(GROUPS, key=lambda g: g[5]):
    key, label, statement, side, sect, order, etypes, keywords = g
    label_ref = f'Master!$B${group_row_of[key]}'
    notes.cell(row=r, column=2, value=f'=IF(Master!$I${group_row_of[key]}="","",'
                                       f'"Note "&Master!$I${group_row_of[key]}&": "&{label_ref})').font = TOTAL_FONT
    notes.cell(row=r, column=3, value=f'=IF(Master!$I${group_row_of[key]}="","",TEXT(Home!$F$10,"DD MMM YYYY"))')
    notes.cell(row=r, column=4, value=f'=IF(Master!$I${group_row_of[key]}="","",TEXT(Home!$F$11,"DD MMM YYYY"))')
    r += 1
    notes.cell(row=r, column=2, value=f'=IF(Master!$I${group_row_of[key]}="","","Total")')
    notes.cell(row=r, column=3, value=amount_formula(label_ref, side, "CY")).number_format = "#,##0.00;(#,##0.00);-"
    notes.cell(row=r, column=4, value=amount_formula(label_ref, side, "PY")).number_format = "#,##0.00;(#,##0.00);-"
    r += 2

wb.save(os.path.join(SCRIPT_DIR, "_wip.xlsx"))
print("Notes sheet built")

# --------------------------------------------------------------- Finalize --
order = ["Home", "TB", "BS", "PL", "Notes", "Master"]
wb._sheets = [wb[name] for name in order]
wb.active = 0
master.sheet_state = "hidden"

OUT = os.path.join(SCRIPT_DIR, "ICAI_Financial_Statement_Tool.xlsx")
wb.save(OUT)
print("Saved", OUT)
